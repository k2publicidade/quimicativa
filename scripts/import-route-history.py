"""Convert the delivery worksheet into the read-only route history used by the app.

Usage: python scripts/import-route-history.py "C:/path/ROTEIRO PADRÃO.xlsx"
"""

from datetime import datetime
import json
from pathlib import Path
import sys

from openpyxl import load_workbook


def number(value):
    return float(value) if value is not None else 0


def main():
    if len(sys.argv) != 2:
        raise SystemExit("Informe o caminho da planilha de roteiros.")
    source = Path(sys.argv[1])
    sheet = load_workbook(source, read_only=True, data_only=True)["Roteiro de Entregas"]
    records = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, max_col=18, values_only=True), 2):
        if not isinstance(row[0], datetime) or not row[2]:
            continue
        records.append({
            "row": row_number,
            "date": row[0].date().isoformat(),
            "vehicle": str(row[1] or "").strip(),
            "customer": str(row[2] or "").strip(),
            "district": str(row[3] or "").strip(),
            "product": str(row[4] or "").strip(),
            "quantity": number(row[5]),
            "packageKg": str(row[6] or "").strip(),
            "weightKg": number(row[7]),
            "company": str(row[8] or "").strip(),
            "pricePerKg": number(row[9]),
            "revenue": number(row[10]),
            "liters": number(row[11]),
            "productCost": number(row[12]),
            "tax": number(row[13]),
            "commission": number(row[14]),
            "grossProfit": number(row[15]),
            "operatingCost": number(row[16]),
            "netProfit": number(row[17]),
        })
    output = Path(__file__).resolve().parents[1] / "app" / "data" / "route-history.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(records, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(f"{len(records)} registros salvos em {output}")


if __name__ == "__main__":
    main()
